"""Shopify CSV/Excel 解析、订单导入、产品标题缓存与匹配。

由 ``appcore.order_analytics`` package 在 PR 1.3 抽出；函数体逐字符
保留，行为不变。``__init__.py`` 通过显式 re-export 把这里的公开符号
带回 ``appcore.order_analytics`` 命名空间。
"""
from __future__ import annotations

import csv
import io
import logging
import sys
import time

import requests

from ._constants import _TITLE_RE
from ._helpers import _parse_shopify_ts, _safe_float, _safe_int

log = logging.getLogger(__name__)


# DB 入口走 module-level wrapper（与 dianxiaomi.py 同样原理）：让现有
# monkeypatch.setattr(oa, "query", fake) 透传到本模块的 query(...) 调用。
def _facade():
    return sys.modules[__package__]


def query(*args, **kwargs):
    return _facade().query(*args, **kwargs)


def query_one(*args, **kwargs):
    return _facade().query_one(*args, **kwargs)


def execute(*args, **kwargs):
    return _facade().execute(*args, **kwargs)


def get_conn(*args, **kwargs):
    return _facade().get_conn(*args, **kwargs)


def parse_shopify_file(file_stream, filename: str) -> list[dict]:
    """解析 CSV 或 Excel 文件，返回原始行 dict 列表。"""
    filename = filename.lower()
    if filename.endswith(".csv"):
        text = file_stream.read()
        if isinstance(text, bytes):
            text = text.decode("utf-8-sig")
        return list(csv.DictReader(io.StringIO(text)))
    elif filename.endswith((".xls", ".xlsx")):
        return _parse_excel(file_stream)
    else:
        raise ValueError("仅支持 CSV / Excel (.xlsx) 文件")


def _parse_excel(stream) -> list[dict]:
    try:
        import openpyxl
    except ImportError:
        raise RuntimeError("服务器未安装 openpyxl，无法解析 Excel 文件")
    wb = openpyxl.load_workbook(stream, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    headers = next(rows_iter)
    if not headers:
        return []
    headers = [str(h).strip() if h else "" for h in headers]
    result = []
    for row in rows_iter:
        d = {}
        for i, val in enumerate(row):
            if i < len(headers) and headers[i]:
                d[headers[i]] = str(val) if val is not None else ""
        result.append(d)
    wb.close()
    return result


def import_orders(rows: list[dict]) -> dict:
    """将原始行批量写入 shopify_orders，去重。返回 {imported, skipped}。"""
    if not rows:
        return {"imported": 0, "skipped": 0}

    BATCH = 500
    imported = 0
    skipped = 0

    for i in range(0, len(rows), BATCH):
        batch = rows[i:i + BATCH]
        values = []
        for row in batch:
            oid_str = (row.get("Id") or "").strip()
            if not oid_str:
                skipped += 1
                continue
            try:
                oid = int(oid_str)
            except ValueError:
                skipped += 1
                continue

            name = (row.get("Lineitem name") or "").strip()
            if not name:
                skipped += 1
                continue

            values.append((
                oid,
                (row.get("Name") or "").strip()[:32] or None,
                _parse_shopify_ts(row.get("Created at", "")),
                name[:500],
                (row.get("Lineitem sku") or "").strip()[:128] or None,
                _safe_int(row.get("Lineitem quantity", "1"), 1),
                _safe_float(row.get("Lineitem price")),
                (row.get("Billing Country") or "").strip()[:8] or None,
                _safe_float(row.get("Total")),
                _safe_float(row.get("Subtotal")),
                _safe_float(row.get("Shipping")),
                (row.get("Currency") or "").strip()[:8] or None,
                (row.get("Financial Status") or "").strip()[:32] or None,
                (row.get("Fulfillment Status") or "").strip()[:32] or None,
                (row.get("Vendor") or "").strip()[:128] or None,
            ))

        if not values:
            continue

        placeholders = ",".join(["(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)"] * len(values))
        flat = []
        for v in values:
            flat.extend(v)

        sql = (
            "INSERT IGNORE INTO shopify_orders "
            "(shopify_order_id, order_name, created_at_order, lineitem_name, "
            "lineitem_sku, lineitem_quantity, lineitem_price, billing_country, "
            "total, subtotal, shipping, currency, financial_status, "
            "fulfillment_status, vendor) VALUES " + placeholders
        )
        conn = get_conn()
        try:
            with conn.cursor() as cur:
                cur.execute(sql, tuple(flat))
                affected = cur.rowcount
        finally:
            conn.close()
        imported += affected
        skipped += len(values) - affected

    return {"imported": imported, "skipped": skipped}


def get_import_stats() -> dict:
    """返回数据库统计概览。"""
    row = query_one(
        "SELECT COUNT(*) AS total_rows, "
        "MIN(created_at_order) AS min_date, MAX(created_at_order) AS max_date, "
        "COUNT(DISTINCT lineitem_name) AS product_count, "
        "COUNT(DISTINCT billing_country) AS country_count, "
        "SUM(CASE WHEN product_id IS NOT NULL THEN 1 ELSE 0 END) AS matched_rows "
        "FROM shopify_orders"
    )
    return row or {}


def fetch_product_page_title(product_code: str) -> str | None:
    """从英文产品页抓取 <title>。"""
    url = f"https://newjoyloo.com/products/{product_code}"
    try:
        resp = requests.get(url, timeout=8, allow_redirects=True,
                            headers={"User-Agent": "Mozilla/5.0 AutoVideoSrt/1.0"})
        if resp.status_code != 200:
            return None
        m = _TITLE_RE.search(resp.text)
        if m:
            title = m.group(1).strip()
            # 先解码 HTML 实体
            title = (title
                     .replace("&ndash;", "–")
                     .replace("&mdash;", "—")
                     .replace("&amp;", "&")
                     .replace("&lt;", "<")
                     .replace("&gt;", ">")
                     .replace("&quot;", '"'))
            # 再去掉 " | Store Name" / " – Store Name" / " - Store Name" 后缀
            for sep in (" | ", " – ", " — ", " - "):
                if sep in title:
                    title = title.rsplit(sep, 1)[0].strip()
                    break
            return title[:500] if title else None
    except requests.RequestException as exc:
        log.debug("fetch title failed for %s: %s", product_code, exc)
    return None


def refresh_product_titles(product_ids: list[int] | None = None) -> dict:
    """批量刷新产品标题缓存。返回 {fetched, skipped, errors}。"""
    if product_ids:
        placeholders = ",".join(["%s"] * len(product_ids))
        products = query(
            f"SELECT id, product_code FROM media_products "
            f"WHERE id IN ({placeholders}) AND product_code IS NOT NULL AND deleted_at IS NULL",
            tuple(product_ids),
        )
    else:
        products = query(
            "SELECT mp.id, mp.product_code FROM media_products mp "
            "LEFT JOIN product_title_cache ptc ON ptc.product_id = mp.id "
            "WHERE mp.product_code IS NOT NULL AND mp.deleted_at IS NULL "
            "AND (ptc.id IS NULL OR ptc.fetched_at < DATE_SUB(NOW(), INTERVAL 7 DAY))",
        )

    fetched = 0
    errors = 0
    for p in products:
        code = p["product_code"]
        if not code:
            continue
        title = fetch_product_page_title(code)
        if title:
            execute(
                "INSERT INTO product_title_cache (product_id, product_code, page_title, fetched_at) "
                "VALUES (%s, %s, %s, NOW()) "
                "ON DUPLICATE KEY UPDATE page_title=VALUES(page_title), fetched_at=NOW()",
                (p["id"], code, title),
            )
            fetched += 1
        else:
            errors += 1
        time.sleep(0.5)  # 限速

    return {"fetched": fetched, "skipped": len(products) - fetched - errors, "errors": errors}


def match_orders_to_products() -> int:
    """将 lineitem_name 匹配到 product_id。返回新匹配行数。"""
    # 前缀匹配：订单商品名以产品标题开头（后面可能有变体信息如 "- 1 Pack"）
    affected = execute(
        "UPDATE shopify_orders so "
        "JOIN product_title_cache ptc ON so.lineitem_name LIKE CONCAT(ptc.page_title, '%') "
        "SET so.product_id = ptc.product_id "
        "WHERE so.product_id IS NULL"
    )
    return affected
