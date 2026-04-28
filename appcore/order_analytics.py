"""订单分析 DAO 层：Shopify 订单导入、产品匹配、数据分析查询。"""
from __future__ import annotations

import calendar
import csv
import hashlib
import io
import json
import logging
import re
import time
from dataclasses import dataclass
from datetime import date, datetime, time as dt_time, timedelta
from typing import Any
from zoneinfo import ZoneInfo

import requests

from appcore.db import query, query_one, execute, get_conn

log = logging.getLogger(__name__)

META_ATTRIBUTION_CUTOVER_HOUR_BJ = 16
META_ATTRIBUTION_TIMEZONE = "Asia/Shanghai"

# Shopify CSV 列名映射
_SHOPIFY_COLS = {
    "Id":                   "shopify_order_id",
    "Name":                 "order_name",
    "Created at":           "created_at_order",
    "Lineitem name":        "lineitem_name",
    "Lineitem sku":         "lineitem_sku",
    "Lineitem quantity":    "lineitem_quantity",
    "Lineitem price":       "lineitem_price",
    "Billing Country":      "billing_country",
    "Total":                "total",
    "Subtotal":             "subtotal",
    "Shipping":             "shipping",
    "Currency":             "currency",
    "Financial Status":     "financial_status",
    "Fulfillment Status":   "fulfillment_status",
    "Vendor":               "vendor",
}

_TITLE_RE = re.compile(r"<title[^>]*>(.*?)</title>", re.IGNORECASE | re.DOTALL)
_SHOP_TS_FMT = "%Y-%m-%d %H:%M:%S %z"  # "2026-04-22 23:00:14 -0700"

_META_AD_REQUIRED_COLS = [
    "报告开始日期",
    "报告结束日期",
    "广告系列名称",
    "已花费金额 (USD)",
]

_META_AD_NUMERIC_FIELDS: dict[str, tuple[str, str]] = {
    "成效": ("result_count", "int"),
    "已花费金额 (USD)": ("spend_usd", "float"),
    "购物转化价值": ("purchase_value_usd", "float"),
    "广告花费回报 (ROAS) - 购物": ("roas_purchase", "float"),
    "CPM（千次展示费用） (USD)": ("cpm_usd", "float"),
    "单次链接点击费用 - 独立用户 (USD)": ("unique_link_click_cost_usd", "float"),
    "链接点击率": ("link_ctr", "float"),
    "链接点击量": ("link_clicks", "int"),
    "加入购物车次数": ("add_to_cart_count", "int"),
    "结账发起次数": ("initiate_checkout_count", "int"),
    "单次加入购物车费用 (USD)": ("add_to_cart_cost_usd", "float"),
    "单次发起结账费用 (USD)": ("initiate_checkout_cost_usd", "float"),
    "单次成效费用": ("cost_per_result_usd", "float"),
    "平均购物转化价值": ("average_purchase_value_usd", "float"),
    "展示次数": ("impressions", "int"),
    "视频平均播放时长": ("video_avg_play_time", "float"),
}

_DIANXIAOMI_SITE_DOMAINS: dict[str, tuple[str, ...]] = {
    "newjoy": ("newjoyloo.com",),
    "omurio": ("omurio.com", "omurio"),
}
_DIANXIAOMI_EXCLUDED_DOMAINS = ("smartgearx.com", "smartgearx")


@dataclass(frozen=True)
class DianxiaomiProductScope:
    by_shopify_id: dict[str, dict[str, Any]]
    by_handle: dict[str, dict[str, Any]]
    excluded_shopify_ids: set[str]
    excluded_handles: set[str]
    requested_site_codes: set[str]


def _safe_decimal_float(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        return round(float(str(value).replace(",", "").strip()), 2)
    except (TypeError, ValueError):
        return None


def _parse_dianxiaomi_ts(value: Any) -> datetime | None:
    if isinstance(value, (int, float)):
        timestamp = float(value)
        if timestamp > 10_000_000_000:
            timestamp = timestamp / 1000
        try:
            return datetime.fromtimestamp(timestamp).replace(microsecond=0)
        except (OSError, OverflowError, ValueError):
            return None
    text = str(value or "").strip()
    if not text:
        return None
    if text.isdigit():
        return _parse_dianxiaomi_ts(int(text))
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        candidate = text[:19] if fmt.endswith("%S") else text[:10]
        try:
            return datetime.strptime(candidate, fmt)
        except ValueError:
            continue
    return None


def compute_meta_business_window_bj(day_value: date) -> tuple[datetime, datetime]:
    start = datetime(day_value.year, day_value.month, day_value.day, META_ATTRIBUTION_CUTOVER_HOUR_BJ, 0, 0)
    return start, start + timedelta(days=1)


def compute_order_meta_attribution(
    *values: datetime | None,
) -> dict[str, Any]:
    attribution_time = next((value for value in values if value is not None), None)
    if attribution_time is None:
        return {
            "attribution_time_at": None,
            "attribution_source": None,
            "attribution_timezone": META_ATTRIBUTION_TIMEZONE,
            "meta_business_date": None,
            "meta_window_start_at": None,
            "meta_window_end_at": None,
        }
    business_date = (attribution_time - timedelta(hours=META_ATTRIBUTION_CUTOVER_HOUR_BJ)).date()
    window_start, window_end = compute_meta_business_window_bj(business_date)
    return {
        "attribution_time_at": attribution_time,
        "attribution_source": None,
        "attribution_timezone": META_ATTRIBUTION_TIMEZONE,
        "meta_business_date": business_date,
        "meta_window_start_at": window_start,
        "meta_window_end_at": window_end,
    }


def _combined_link_text(*values: Any) -> str:
    return " ".join(str(value or "") for value in values).lower()


def _infer_dianxiaomi_site_code_from_text(text: str, requested_site_codes: set[str]) -> str | None:
    normalized = (text or "").lower()
    if not normalized:
        return None
    if any(domain in normalized for domain in _DIANXIAOMI_EXCLUDED_DOMAINS):
        return "smartgearx"
    for site_code in sorted(requested_site_codes):
        domains = _DIANXIAOMI_SITE_DOMAINS.get(site_code, (site_code,))
        if any(domain in normalized for domain in domains):
            return site_code
    return None


def extract_dianxiaomi_shopify_product_id(line: dict[str, Any]) -> str | None:
    for key in ("productId", "shopifyProductId", "pid"):
        value = str(line.get(key) or "").strip()
        if value.isdigit():
            return value
    for key in ("productUrl", "sourceUrl"):
        text = str(line.get(key) or "")
        match = re.search(r"/products/(\d+)", text)
        if match:
            return match.group(1)
    return None


def _canonical_product_handle(value: Any) -> str | None:
    text = str(value or "").strip().lower()
    if not text:
        return None
    if "/products/" in text:
        text = text.split("/products/", 1)[1]
    text = text.split("?", 1)[0].split("#", 1)[0].strip("/")
    if not text:
        return None
    if text.endswith("-rjc"):
        text = text[:-4]
    return text or None


def extract_dianxiaomi_product_handle(line: dict[str, Any]) -> str | None:
    for key in ("productUrl", "sourceUrl"):
        handle = _canonical_product_handle(line.get(key))
        if handle:
            return handle
    return None


def build_dianxiaomi_product_scope(site_codes: list[str]) -> DianxiaomiProductScope:
    requested = {str(code).strip().lower() for code in site_codes if str(code).strip()}
    rows = query(
        "SELECT id, product_code, shopifyid, product_link, localized_links_json "
        "FROM media_products "
        "WHERE deleted_at IS NULL AND shopifyid IS NOT NULL AND shopifyid <> ''"
    )
    by_shopify_id: dict[str, dict[str, Any]] = {}
    by_handle: dict[str, dict[str, Any]] = {}
    excluded_shopify_ids: set[str] = set()
    excluded_handles: set[str] = set()
    for row in rows:
        shopifyid = str(row.get("shopifyid") or "").strip()
        product_code = str(row.get("product_code") or "").strip()
        handle = _canonical_product_handle(product_code)
        if not shopifyid and not handle:
            continue
        site_code = _infer_dianxiaomi_site_code_from_text(
            _combined_link_text(
                product_code,
                row.get("product_link"),
                row.get("localized_links_json"),
            ),
            requested,
        )
        if site_code == "smartgearx":
            if shopifyid:
                excluded_shopify_ids.add(shopifyid)
            if handle:
                excluded_handles.add(handle)
            continue
        if site_code in requested:
            product = {
                "product_id": row.get("id"),
                "product_code": product_code,
                "site_code": site_code,
                "shopifyid": shopifyid,
            }
            if shopifyid:
                by_shopify_id[shopifyid] = product
            if handle:
                by_handle[handle] = product
    return DianxiaomiProductScope(
        by_shopify_id=by_shopify_id,
        by_handle=by_handle,
        excluded_shopify_ids=excluded_shopify_ids,
        excluded_handles=excluded_handles,
        requested_site_codes=requested,
    )


def _dianxiaomi_order_lines(order: dict[str, Any]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for key in ("productList", "cancelProductList"):
        value = order.get(key) or []
        if isinstance(value, list):
            rows.extend(item for item in value if isinstance(item, dict))
    return rows


def _resolve_dianxiaomi_line_product(
    line: dict[str, Any],
    shopify_product_id: str,
    scope: DianxiaomiProductScope,
) -> dict[str, Any] | None:
    handle = extract_dianxiaomi_product_handle(line)
    if handle in scope.excluded_handles:
        return None
    if shopify_product_id in scope.excluded_shopify_ids:
        return None
    product = scope.by_shopify_id.get(shopify_product_id)
    if product:
        return product
    if handle:
        product = scope.by_handle.get(handle)
        if product:
            return product
    site_code = _infer_dianxiaomi_site_code_from_text(
        _combined_link_text(line.get("productUrl"), line.get("sourceUrl")),
        scope.requested_site_codes,
    )
    if site_code == "smartgearx" or site_code not in scope.requested_site_codes:
        return None
    return {
        "product_id": None,
        "product_code": None,
        "site_code": site_code,
        "shopifyid": shopify_product_id,
    }


def normalize_dianxiaomi_order(
    order: dict[str, Any],
    scope: DianxiaomiProductScope,
    profits_by_package_id: dict[str, dict[str, Any]],
) -> tuple[list[dict[str, Any]], int]:
    normalized: list[dict[str, Any]] = []
    skipped = 0
    package_id = str(order.get("id") or order.get("packageId") or "").strip()
    profit = profits_by_package_id.get(package_id) or {}
    for line in _dianxiaomi_order_lines(order):
        shopify_product_id = extract_dianxiaomi_shopify_product_id(line)
        if not shopify_product_id:
            skipped += 1
            continue
        product = _resolve_dianxiaomi_line_product(line, shopify_product_id, scope)
        if not product:
            skipped += 1
            continue
        quantity = _safe_int(str(line.get("quantity") or line.get("productCount") or "1"), 1)
        unit_price = _safe_decimal_float(line.get("price"))
        line_amount = round((unit_price or 0) * quantity, 2) if unit_price is not None else None
        addr = order.get("dxmPackageAddr") if isinstance(order.get("dxmPackageAddr"), dict) else {}
        order_created_at = _parse_dianxiaomi_ts(order.get("orderCreateTime"))
        order_paid_at = _parse_dianxiaomi_ts(order.get("orderPayTime"))
        paid_at = _parse_dianxiaomi_ts(order.get("paidTime"))
        shipped_at = _parse_dianxiaomi_ts(order.get("shippedTime"))
        attribution = compute_order_meta_attribution(order_paid_at, paid_at, order_created_at, shipped_at)
        attribution["attribution_source"] = (
            "order_paid_at" if order_paid_at is not None else
            "paid_at" if paid_at is not None else
            "order_created_at" if order_created_at is not None else
            "shipped_at" if shipped_at is not None else None
        )
        normalized.append({
            "site_code": product["site_code"],
            "product_id": product["product_id"],
            "product_code": product["product_code"],
            "shopify_product_id": shopify_product_id,
            "dxm_shop_id": str(order.get("shopId") or "").strip() or None,
            "dxm_shop_name": str(order.get("shopName") or "").strip() or None,
            "dxm_package_id": package_id,
            "dxm_order_id": str(order.get("orderId") or "").strip() or None,
            "extended_order_id": str(order.get("extendedOrderId") or "").strip() or None,
            "package_number": str(order.get("packageNumber") or "").strip() or None,
            "platform": str(order.get("platform") or order.get("shopPlatform") or "").strip() or None,
            "order_state": str(order.get("state") or "").strip() or None,
            "buyer_name": str(order.get("buyerName") or "").strip() or None,
            "buyer_account": str(order.get("buyerAccount") or "").strip() or None,
            "product_name": str(line.get("productName") or "").strip()[:500] or None,
            "product_sku": str(line.get("productSku") or "").strip()[:128] or None,
            "product_sub_sku": str(line.get("productSubSku") or "").strip()[:128] or None,
            "product_display_sku": str(line.get("productDisplaySku") or line.get("displaySku") or "").strip()[:128] or None,
            "variant_text": str(line.get("attrListStr") or line.get("attrList") or "").strip()[:500] or None,
            "quantity": quantity,
            "unit_price": unit_price,
            "line_amount": line_amount,
            "order_amount": _safe_decimal_float(order.get("orderAmount")),
            "order_currency": str(order.get("orderUnit") or "").strip() or None,
            "ship_amount": _safe_decimal_float(order.get("shipAmount")),
            "amount_with_shipping": _safe_decimal_float(order.get("orderAmount")),
            "amount_cny": _safe_decimal_float(profit.get("amountCNY")),
            "logistic_fee": _safe_decimal_float(profit.get("logisticFee")),
            "profit": _safe_decimal_float(profit.get("profit")),
            "refund_amount_usd": _safe_decimal_float(order.get("refundAmountUsd")),
            "refund_amount": _safe_decimal_float(order.get("refundAmount")),
            "buyer_country": str(order.get("buyerCountry") or addr.get("country") or "").strip() or None,
            "buyer_country_name": str(order.get("countryCN") or "").strip() or None,
            "province": str(addr.get("province") or "").strip() or None,
            "city": str(addr.get("city") or "").strip() or None,
            "order_created_at": order_created_at,
            "order_paid_at": order_paid_at,
            "paid_at": paid_at,
            "shipped_at": shipped_at,
            **attribution,
            "raw_order_json": order,
            "raw_line_json": line,
            "profit_json": profit or None,
        })
    return normalized, skipped


def start_dianxiaomi_order_import_batch(
    date_from: str,
    date_to: str,
    site_codes: list[str],
    included_shopify_ids_count: int,
) -> int:
    conn = get_conn()
    try:
        with conn.cursor() as cur:
            cur.execute(
                "INSERT INTO dianxiaomi_order_import_batches "
                "(date_from, date_to, requested_site_codes, included_shopify_ids_count) "
                "VALUES (%s,%s,%s,%s)",
                (date_from, date_to, ",".join(site_codes), included_shopify_ids_count),
            )
            batch_id = int(cur.lastrowid)
        conn.commit()
        return batch_id
    finally:
        conn.close()


def finish_dianxiaomi_order_import_batch(
    batch_id: int,
    status: str,
    summary: dict[str, Any],
    error_message: str | None = None,
) -> None:
    execute(
        "UPDATE dianxiaomi_order_import_batches SET status=%s, finished_at=NOW(), "
        "duration_seconds=TIMESTAMPDIFF(SECOND, started_at, NOW()), "
        "total_pages=%s, fetched_orders=%s, fetched_lines=%s, inserted_lines=%s, "
        "updated_lines=%s, skipped_lines=%s, error_message=%s, summary_json=%s "
        "WHERE id=%s",
        (
            status,
            int(summary.get("total_pages") or 0),
            int(summary.get("fetched_orders") or 0),
            int(summary.get("fetched_lines") or 0),
            int(summary.get("inserted_lines") or 0),
            int(summary.get("updated_lines") or 0),
            int(summary.get("skipped_lines") or 0),
            error_message,
            json.dumps(summary, ensure_ascii=False, default=str),
            batch_id,
        ),
    )


def _json_dumps_for_db(value: Any) -> str | None:
    if value is None:
        return None
    return json.dumps(value, ensure_ascii=False, default=str)


_DIANXIAOMI_ORDER_LINE_COLUMNS = [
    "batch_id",
    "site_code",
    "product_id",
    "product_code",
    "shopify_product_id",
    "dxm_shop_id",
    "dxm_shop_name",
    "dxm_package_id",
    "dxm_order_id",
    "extended_order_id",
    "package_number",
    "platform",
    "order_state",
    "buyer_name",
    "buyer_account",
    "product_name",
    "product_sku",
    "product_sub_sku",
    "product_display_sku",
    "variant_text",
    "quantity",
    "unit_price",
    "line_amount",
    "order_amount",
    "order_currency",
    "ship_amount",
    "amount_with_shipping",
    "amount_cny",
    "logistic_fee",
    "profit",
    "refund_amount_usd",
    "refund_amount",
    "buyer_country",
    "buyer_country_name",
    "province",
    "city",
    "order_created_at",
    "order_paid_at",
    "paid_at",
    "shipped_at",
    "attribution_time_at",
    "attribution_source",
    "attribution_timezone",
    "meta_business_date",
    "meta_window_start_at",
    "meta_window_end_at",
    "raw_order_json",
    "raw_line_json",
    "profit_json",
]


def _dianxiaomi_order_line_values(batch_id: int, row: dict[str, Any]) -> tuple[Any, ...]:
    enriched = dict(row)
    enriched["batch_id"] = batch_id
    enriched["raw_order_json"] = _json_dumps_for_db(row.get("raw_order_json"))
    enriched["raw_line_json"] = _json_dumps_for_db(row.get("raw_line_json"))
    enriched["profit_json"] = _json_dumps_for_db(row.get("profit_json"))
    return tuple(enriched.get(column) for column in _DIANXIAOMI_ORDER_LINE_COLUMNS)


def upsert_dianxiaomi_order_lines(batch_id: int, rows: list[dict[str, Any]]) -> dict[str, int]:
    if not rows:
        return {"affected": 0, "rows": 0}
    columns_sql = ", ".join(_DIANXIAOMI_ORDER_LINE_COLUMNS)
    placeholders = ", ".join(["%s"] * len(_DIANXIAOMI_ORDER_LINE_COLUMNS))
    update_columns = [
        "batch_id",
        "site_code",
        "product_id",
        "product_code",
        "quantity",
        "unit_price",
        "line_amount",
        "order_amount",
        "order_currency",
        "ship_amount",
        "amount_with_shipping",
        "amount_cny",
        "logistic_fee",
        "profit",
        "refund_amount_usd",
        "refund_amount",
        "buyer_country",
        "buyer_country_name",
        "province",
        "city",
        "order_created_at",
        "order_paid_at",
        "paid_at",
        "shipped_at",
        "attribution_time_at",
        "attribution_source",
        "attribution_timezone",
        "meta_business_date",
        "meta_window_start_at",
        "meta_window_end_at",
        "raw_order_json",
        "raw_line_json",
        "profit_json",
    ]
    updates_sql = ", ".join(f"{column}=VALUES({column})" for column in update_columns)
    sql = (
        f"INSERT INTO dianxiaomi_order_lines ({columns_sql}) "
        f"VALUES ({placeholders}) "
        f"ON DUPLICATE KEY UPDATE {updates_sql}, imported_at=NOW()"
    )

    affected = 0
    conn = get_conn()
    try:
        with conn.cursor() as cur:
            for row in rows:
                cur.execute(sql, _dianxiaomi_order_line_values(batch_id, row))
                affected += int(cur.rowcount or 0)
        conn.commit()
    finally:
        conn.close()
    return {"affected": affected, "rows": len(rows)}


def get_dianxiaomi_order_import_batches(limit: int = 20) -> list[dict]:
    limit = max(1, min(int(limit or 20), 100))
    return query(
        "SELECT * FROM dianxiaomi_order_import_batches "
        "ORDER BY started_at DESC LIMIT %s",
        (limit,),
    )


# ── 解析 ───────────────────────────────────────────────

def parse_shopify_file(file_stream, filename: str) -> list[dict]:
    """解析 CSV 或 Excel 文件，返回原始行 dict 列表。"""
    filename = filename.lower()
    if filename.endswith(".csv"):
        text = file_stream.read()
        if isinstance(text, bytes):
            text = text.decode("utf-8-sig")
        return list(csv.DictReader(io.StringIO(text)))
    elif filename.endswith((".xls", ".xlsx")):
        return _parse_excel(file_stream)
    else:
        raise ValueError("仅支持 CSV / Excel (.xlsx) 文件")


def _parse_excel(stream) -> list[dict]:
    try:
        import openpyxl
    except ImportError:
        raise RuntimeError("服务器未安装 openpyxl，无法解析 Excel 文件")
    wb = openpyxl.load_workbook(stream, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    headers = next(rows_iter)
    if not headers:
        return []
    headers = [str(h).strip() if h else "" for h in headers]
    result = []
    for row in rows_iter:
        d = {}
        for i, val in enumerate(row):
            if i < len(headers) and headers[i]:
                d[headers[i]] = str(val) if val is not None else ""
        result.append(d)
    wb.close()
    return result


def _parse_shopify_ts(ts_str: str) -> datetime | None:
    """解析 Shopify 时间戳 '2026-04-22 23:00:14 -0700' 为 naive UTC-ish datetime。"""
    ts_str = (ts_str or "").strip()
    if not ts_str:
        return None
    try:
        dt = datetime.strptime(ts_str, _SHOP_TS_FMT)
        # 转为 UTC（去掉时区信息）
        return dt.replace(tzinfo=None) - dt.utcoffset()
    except Exception:
        pass
    # fallback: 只取日期时间部分
    try:
        return datetime.strptime(ts_str[:19], "%Y-%m-%d %H:%M:%S")
    except Exception:
        return None


def _safe_int(val: str, default: int = 0) -> int:
    try:
        return int(float((val or "").strip()))
    except (ValueError, TypeError):
        return default


def _safe_float(val: str) -> float | None:
    try:
        return float((val or "").strip())
    except (ValueError, TypeError):
        return None


def _safe_float_default(val: str, default: float = 0.0) -> float:
    parsed = _safe_float(val)
    return default if parsed is None else parsed


def _parse_meta_date(value: str) -> date:
    value = (value or "").strip()
    if not value:
        raise ValueError("Meta 广告报表日期不能为空")
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%m/%d/%Y"):
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"无法解析 Meta 广告报表日期：{value}")


def product_code_candidates_for_ad_campaign(campaign_name: str) -> list[str]:
    code = (campaign_name or "").strip().lower()
    if not code:
        return []
    candidates = [code]
    if code.endswith("-rjc"):
        candidates.append(code[:-4])
    else:
        candidates.append(f"{code}-rjc")
    return list(dict.fromkeys(candidates))


def resolve_ad_product_match(campaign_name: str) -> dict | None:
    for code in product_code_candidates_for_ad_campaign(campaign_name):
        product = query_one(
            "SELECT id, product_code, name FROM media_products "
            "WHERE product_code=%s AND deleted_at IS NULL",
            (code,),
        )
        if product:
            return product
    return None


def _normalize_meta_ad_row(row: dict) -> dict:
    campaign_name = (row.get("广告系列名称") or "").strip()
    normalized = campaign_name.lower()
    out = {
        "report_start_date": _parse_meta_date(row.get("报告开始日期", "")),
        "report_end_date": _parse_meta_date(row.get("报告结束日期", "")),
        "campaign_name": campaign_name,
        "normalized_campaign_code": normalized,
        "result_metric": (row.get("成效指标") or "").strip()[:128] or None,
        "campaign_delivery": (row.get("广告系列投放") or "").strip()[:32] or None,
        "raw": dict(row),
    }
    for source_col, (target_col, kind) in _META_AD_NUMERIC_FIELDS.items():
        if kind == "int":
            out[target_col] = _safe_int(row.get(source_col, "0"), 0)
        else:
            out[target_col] = _safe_float_default(row.get(source_col, ""), 0.0)
    return out


def parse_meta_ad_file(file_stream, filename: str) -> list[dict]:
    """解析 Meta 广告后台导出的 CSV/Excel，返回标准化后的周期广告行。"""
    rows = parse_shopify_file(file_stream, filename)
    if not rows:
        return []
    headers = set(rows[0].keys())
    missing = [col for col in _META_AD_REQUIRED_COLS if col not in headers]
    if missing:
        raise ValueError("Meta 广告报表缺少列：" + "、".join(missing))
    return [
        _normalize_meta_ad_row(row)
        for row in rows
        if (row.get("广告系列名称") or "").strip()
    ]


def _parse_iso_date_param(value: str, name: str) -> date:
    try:
        return datetime.strptime((value or "").strip(), "%Y-%m-%d").date()
    except ValueError as exc:
        raise ValueError(f"{name} must be YYYY-MM-DD") from exc


def _money(value: Any) -> float:
    return round(float(value or 0), 2)


def _roas(revenue: float, spend: float) -> float | None:
    if spend <= 0:
        return None
    return round(revenue / spend, 4)


def _beijing_now() -> datetime:
    return datetime.now(ZoneInfo(META_ATTRIBUTION_TIMEZONE)).replace(tzinfo=None)


def get_realtime_roas_overview(date_text: str | None = None, now: datetime | None = None) -> dict:
    now = (now or _beijing_now()).replace(microsecond=0)
    target = _parse_iso_date_param(date_text, "date") if date_text else now.date()
    day_start = datetime.combine(target, dt_time.min)
    day_end = day_start + timedelta(days=1)
    if target == now.date():
        data_until = min(now, day_end)
        complete_hour_until = now.replace(minute=0, second=0, microsecond=0)
    elif target < now.date():
        data_until = day_end
        complete_hour_until = day_end
    else:
        data_until = day_start
        complete_hour_until = day_start

    latest_snapshot = query(
        "SELECT * FROM roi_realtime_daily_snapshots "
        "WHERE business_date=%s AND store_scope='newjoy,omurio' AND ad_platform_scope='meta' "
        "ORDER BY snapshot_at DESC, id DESC LIMIT 1",
        (target,),
    )
    if latest_snapshot:
        snap = latest_snapshot[0]
        order_revenue = _money(snap.get("order_revenue_usd"))
        ad_spend = _money(snap.get("ad_spend_usd"))
        return {
            "period": {
                "date": target,
                "timezone": META_ATTRIBUTION_TIMEZONE,
                "day_start_at": day_start,
                "day_end_at": day_end,
                "data_until_at": snap.get("snapshot_at") or data_until,
                "complete_hour_until_at": complete_hour_until,
                "meta_cutover_hour_bj": META_ATTRIBUTION_CUTOVER_HOUR_BJ,
            },
            "scope": {
                "stores": ["newjoy", "omurio"],
                "ad_platforms": ["meta"],
                "order_source": "dianxiaomi",
                "ad_source": "roi_realtime_daily_snapshots",
                "ad_granularity": "day_realtime_snapshot",
                "hourly_ad_ready": False,
            },
            "freshness": {
                "first_order_at": None,
                "last_order_at": snap.get("last_order_at"),
            },
            "summary": {
                "order_count": int(snap.get("order_count") or 0),
                "line_count": int(snap.get("line_count") or 0),
                "units": int(snap.get("units") or 0),
                "order_revenue": order_revenue,
                "line_revenue": 0.0,
                "shipping_revenue": _money(snap.get("shipping_revenue_usd")),
                "ad_spend": ad_spend,
                "meta_purchase_value": 0.0,
                "meta_purchases": 0,
                "true_roas": _roas(order_revenue, ad_spend),
                "order_data_status": snap.get("order_data_status") or "ok",
                "ad_data_status": snap.get("ad_data_status") or "pending_source",
            },
            "hourly": [],
            "snapshots": [snap],
        }

    order_time_expr = "COALESCE(order_paid_at, attribution_time_at, order_created_at)"
    order_rows = query(
        "SELECT HOUR(" + order_time_expr + ") AS hour, "
        "COUNT(DISTINCT dxm_package_id) AS order_count, "
        "COUNT(*) AS line_count, "
        "SUM(quantity) AS units, "
        "SUM(COALESCE(amount_with_shipping, line_amount, 0)) AS order_revenue, "
        "SUM(COALESCE(line_amount, 0)) AS line_revenue, "
        "SUM(COALESCE(ship_amount, 0)) AS shipping_revenue, "
        "MIN(" + order_time_expr + ") AS first_order_at, "
        "MAX(" + order_time_expr + ") AS last_order_at "
        "FROM dianxiaomi_order_lines "
        "WHERE site_code IN ('newjoy', 'omurio') "
        "AND " + order_time_expr + " >= %s AND " + order_time_expr + " < %s "
        "GROUP BY HOUR(" + order_time_expr + ") "
        "ORDER BY hour",
        (day_start, day_end),
    )
    ad_rows = query(
        "SELECT SUM(spend_usd) AS ad_spend, "
        "SUM(purchase_value_usd) AS meta_purchase_value, "
        "SUM(result_count) AS meta_purchases "
        "FROM meta_ad_daily_campaign_metrics "
        "WHERE meta_business_date = %s",
        (target,),
    )

    orders_by_hour = {int(row["hour"]): row for row in order_rows if row.get("hour") is not None}
    ad = ad_rows[0] if ad_rows else {}
    summary = {
        "order_count": 0,
        "line_count": 0,
        "units": 0,
        "order_revenue": 0.0,
        "line_revenue": 0.0,
        "shipping_revenue": 0.0,
        "ad_spend": _money(ad.get("ad_spend")),
        "meta_purchase_value": _money(ad.get("meta_purchase_value")),
        "meta_purchases": int(ad.get("meta_purchases") or 0),
    }
    first_order_at = None
    last_order_at = None
    hourly: list[dict[str, Any]] = []
    for hour in range(24):
        row = orders_by_hour.get(hour, {})
        order_revenue = _money(row.get("order_revenue"))
        item = {
            "hour": hour,
            "window_start_at": day_start + timedelta(hours=hour),
            "window_end_at": day_start + timedelta(hours=hour + 1),
            "order_count": int(row.get("order_count") or 0),
            "line_count": int(row.get("line_count") or 0),
            "units": int(row.get("units") or 0),
            "order_revenue": order_revenue,
            "line_revenue": _money(row.get("line_revenue")),
            "shipping_revenue": _money(row.get("shipping_revenue")),
            "ad_spend": None,
            "true_roas": None,
        }
        hourly.append(item)
        for key in ("order_count", "line_count", "units"):
            summary[key] += item[key]
        for key in ("order_revenue", "line_revenue", "shipping_revenue"):
            summary[key] = round(summary[key] + float(item[key]), 2)
        if row.get("first_order_at") and (first_order_at is None or row["first_order_at"] < first_order_at):
            first_order_at = row["first_order_at"]
        if row.get("last_order_at") and (last_order_at is None or row["last_order_at"] > last_order_at):
            last_order_at = row["last_order_at"]

    summary["true_roas"] = _roas(summary["order_revenue"], summary["ad_spend"])
    return {
        "period": {
            "date": target,
            "timezone": META_ATTRIBUTION_TIMEZONE,
            "day_start_at": day_start,
            "day_end_at": day_end,
            "data_until_at": data_until,
            "complete_hour_until_at": complete_hour_until,
            "meta_cutover_hour_bj": META_ATTRIBUTION_CUTOVER_HOUR_BJ,
        },
        "scope": {
            "stores": ["newjoy", "omurio"],
            "ad_platforms": ["meta"],
            "order_source": "dianxiaomi",
            "ad_source": "meta_ad_daily_campaign_metrics",
            "ad_granularity": "daily",
            "hourly_ad_ready": False,
        },
        "freshness": {
            "first_order_at": first_order_at,
            "last_order_at": last_order_at,
        },
        "summary": summary,
        "hourly": hourly,
    }


def get_true_roas_summary(start_date: str, end_date: str) -> dict:
    start = _parse_iso_date_param(start_date, "start_date")
    end = _parse_iso_date_param(end_date, "end_date")
    if end < start:
        raise ValueError("end_date must be >= start_date")

    order_rows = query(
        "SELECT meta_business_date, "
        "COUNT(DISTINCT dxm_package_id) AS order_count, "
        "COUNT(*) AS line_count, "
        "SUM(quantity) AS units, "
        "SUM(COALESCE(amount_with_shipping, line_amount, 0)) AS order_revenue, "
        "SUM(COALESCE(line_amount, 0)) AS line_revenue, "
        "SUM(COALESCE(ship_amount, 0)) AS shipping_revenue "
        "FROM dianxiaomi_order_lines "
        "WHERE meta_business_date >= %s AND meta_business_date <= %s "
        "GROUP BY meta_business_date",
        (start, end),
    )
    ad_rows = query(
        "SELECT meta_business_date, "
        "SUM(spend_usd) AS ad_spend, "
        "SUM(purchase_value_usd) AS meta_purchase_value, "
        "SUM(result_count) AS meta_purchases "
        "FROM meta_ad_daily_campaign_metrics "
        "WHERE meta_business_date >= %s AND meta_business_date <= %s "
        "GROUP BY meta_business_date",
        (start, end),
    )

    orders_by_day = {row["meta_business_date"]: row for row in order_rows}
    ads_by_day = {row["meta_business_date"]: row for row in ad_rows}
    rows: list[dict[str, Any]] = []
    totals = {
        "order_count": 0,
        "line_count": 0,
        "units": 0,
        "order_revenue": 0.0,
        "line_revenue": 0.0,
        "shipping_revenue": 0.0,
        "ad_spend": 0.0,
        "meta_purchase_value": 0.0,
        "meta_purchases": 0,
    }

    current = start
    while current <= end:
        order = orders_by_day.get(current, {})
        ad = ads_by_day.get(current, {})
        window_start, window_end = compute_meta_business_window_bj(current)
        order_revenue = _money(order.get("order_revenue"))
        ad_spend = _money(ad.get("ad_spend"))
        item = {
            "meta_business_date": current,
            "window_start_at": window_start,
            "window_end_at": window_end,
            "order_count": int(order.get("order_count") or 0),
            "line_count": int(order.get("line_count") or 0),
            "units": int(order.get("units") or 0),
            "order_revenue": order_revenue,
            "line_revenue": _money(order.get("line_revenue")),
            "shipping_revenue": _money(order.get("shipping_revenue")),
            "ad_spend": ad_spend,
            "true_roas": _roas(order_revenue, ad_spend),
            "meta_purchase_value": _money(ad.get("meta_purchase_value")),
            "meta_purchases": int(ad.get("meta_purchases") or 0),
        }
        rows.append(item)
        for key in totals:
            totals[key] += item[key]
        current += timedelta(days=1)

    for key in ("order_revenue", "line_revenue", "shipping_revenue", "ad_spend", "meta_purchase_value"):
        totals[key] = round(float(totals[key]), 2)
    summary = dict(totals)
    summary["true_roas"] = _roas(summary["order_revenue"], summary["ad_spend"])
    return {
        "period": {
            "start": start,
            "end": end,
            "timezone": META_ATTRIBUTION_TIMEZONE,
            "cutover_hour_bj": META_ATTRIBUTION_CUTOVER_HOUR_BJ,
        },
        "summary": summary,
        "rows": rows,
    }


def _coerce_ad_frequency(value: str | None) -> str:
    normalized = (value or "custom").strip().lower()
    if normalized not in {"weekly", "monthly", "custom"}:
        return "custom"
    return normalized


def import_meta_ad_rows(
    rows: list[dict],
    filename: str,
    file_bytes: bytes,
    import_frequency: str = "custom",
) -> dict:
    """将 Meta 广告周期报表 upsert 到长期表。"""
    frequency = _coerce_ad_frequency(import_frequency)
    file_hash = hashlib.sha256(file_bytes or b"").hexdigest()
    starts = [row["report_start_date"] for row in rows if row.get("report_start_date")]
    ends = [row["report_end_date"] for row in rows if row.get("report_end_date")]
    report_start = min(starts) if starts else None
    report_end = max(ends) if ends else None

    conn = get_conn()
    imported = 0
    updated = 0
    skipped = 0
    matched = 0
    try:
        with conn.cursor() as cur:
            cur.execute(
                "INSERT INTO meta_ad_import_batches "
                "(source_filename, file_sha256, import_frequency, report_start_date, report_end_date, raw_row_count) "
                "VALUES (%s,%s,%s,%s,%s,%s)",
                (filename, file_hash, frequency, report_start, report_end, len(rows)),
            )
            batch_id = cur.lastrowid

            for row in rows:
                product = resolve_ad_product_match(row["campaign_name"])
                product_id = product.get("id") if product else None
                matched_product_code = product.get("product_code") if product else None
                if product_id:
                    matched += 1
                meta_business_date = None
                meta_window_start_at = None
                meta_window_end_at = None
                if row["report_start_date"] == row["report_end_date"]:
                    meta_business_date = row["report_start_date"]
                    meta_window_start_at, meta_window_end_at = compute_meta_business_window_bj(meta_business_date)
                args = (
                    batch_id,
                    row["report_start_date"],
                    row["report_end_date"],
                    frequency,
                    row["campaign_name"],
                    row["normalized_campaign_code"],
                    matched_product_code,
                    product_id,
                    row.get("result_count") or 0,
                    row.get("result_metric"),
                    row.get("spend_usd") or 0,
                    row.get("purchase_value_usd") or 0,
                    row.get("roas_purchase"),
                    row.get("cpm_usd"),
                    row.get("unique_link_click_cost_usd"),
                    row.get("link_ctr"),
                    row.get("campaign_delivery"),
                    row.get("link_clicks") or 0,
                    row.get("add_to_cart_count") or 0,
                    row.get("initiate_checkout_count") or 0,
                    row.get("add_to_cart_cost_usd"),
                    row.get("initiate_checkout_cost_usd"),
                    row.get("cost_per_result_usd"),
                    row.get("average_purchase_value_usd"),
                    row.get("impressions") or 0,
                    row.get("video_avg_play_time"),
                    json.dumps(row.get("raw") or {}, ensure_ascii=False),
                    meta_business_date,
                    meta_window_start_at,
                    meta_window_end_at,
                    META_ATTRIBUTION_TIMEZONE,
                )
                cur.execute(
                    "INSERT INTO meta_ad_campaign_metrics "
                    "(import_batch_id, report_start_date, report_end_date, import_frequency, "
                    "campaign_name, normalized_campaign_code, matched_product_code, product_id, "
                    "result_count, result_metric, spend_usd, purchase_value_usd, roas_purchase, "
                    "cpm_usd, unique_link_click_cost_usd, link_ctr, campaign_delivery, link_clicks, "
                    "add_to_cart_count, initiate_checkout_count, add_to_cart_cost_usd, "
                    "initiate_checkout_cost_usd, cost_per_result_usd, average_purchase_value_usd, "
                    "impressions, video_avg_play_time, raw_json, "
                    "meta_business_date, meta_window_start_at, meta_window_end_at, attribution_timezone) "
                    "VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s) "
                    "ON DUPLICATE KEY UPDATE "
                    "import_batch_id=VALUES(import_batch_id), import_frequency=VALUES(import_frequency), "
                    "normalized_campaign_code=VALUES(normalized_campaign_code), "
                    "matched_product_code=VALUES(matched_product_code), product_id=VALUES(product_id), "
                    "result_count=VALUES(result_count), result_metric=VALUES(result_metric), "
                    "spend_usd=VALUES(spend_usd), purchase_value_usd=VALUES(purchase_value_usd), "
                    "roas_purchase=VALUES(roas_purchase), cpm_usd=VALUES(cpm_usd), "
                    "unique_link_click_cost_usd=VALUES(unique_link_click_cost_usd), "
                    "link_ctr=VALUES(link_ctr), campaign_delivery=VALUES(campaign_delivery), "
                    "link_clicks=VALUES(link_clicks), add_to_cart_count=VALUES(add_to_cart_count), "
                    "initiate_checkout_count=VALUES(initiate_checkout_count), "
                    "add_to_cart_cost_usd=VALUES(add_to_cart_cost_usd), "
                    "initiate_checkout_cost_usd=VALUES(initiate_checkout_cost_usd), "
                    "cost_per_result_usd=VALUES(cost_per_result_usd), "
                    "average_purchase_value_usd=VALUES(average_purchase_value_usd), "
                    "impressions=VALUES(impressions), video_avg_play_time=VALUES(video_avg_play_time), "
                    "raw_json=VALUES(raw_json), meta_business_date=VALUES(meta_business_date), "
                    "meta_window_start_at=VALUES(meta_window_start_at), "
                    "meta_window_end_at=VALUES(meta_window_end_at), "
                    "attribution_timezone=VALUES(attribution_timezone)",
                    args,
                )
                if cur.rowcount == 1:
                    imported += 1
                elif cur.rowcount == 2:
                    updated += 1
                else:
                    skipped += 1

            cur.execute(
                "UPDATE meta_ad_import_batches SET imported_rows=%s, updated_rows=%s, "
                "skipped_rows=%s, matched_rows=%s WHERE id=%s",
                (imported, updated, skipped, matched, batch_id),
            )
    finally:
        conn.close()

    return {
        "batch_id": batch_id,
        "imported": imported,
        "updated": updated,
        "skipped": skipped,
        "matched": matched,
    }


def match_meta_ads_to_products() -> int:
    rows = query(
        "SELECT id, campaign_name FROM meta_ad_campaign_metrics "
        "WHERE product_id IS NULL OR matched_product_code IS NULL",
    )
    affected = 0
    for row in rows:
        product = resolve_ad_product_match(row.get("campaign_name") or "")
        if not product:
            continue
        affected += execute(
            "UPDATE meta_ad_campaign_metrics SET product_id=%s, matched_product_code=%s WHERE id=%s",
            (product["id"], product["product_code"], row["id"]),
        )
    return affected


def get_meta_ad_stats() -> dict:
    row = query_one(
        "SELECT COUNT(*) AS total_rows, "
        "COUNT(DISTINCT CONCAT(report_start_date, '|', report_end_date)) AS period_count, "
        "MIN(report_start_date) AS min_date, MAX(report_end_date) AS max_date, "
        "SUM(CASE WHEN product_id IS NOT NULL THEN 1 ELSE 0 END) AS matched_rows, "
        "SUM(spend_usd) AS total_spend_usd, "
        "SUM(purchase_value_usd) AS total_purchase_value_usd "
        "FROM meta_ad_campaign_metrics"
    )
    return row or {}


def get_meta_ad_periods() -> list[dict]:
    return query(
        "SELECT MAX(import_batch_id) AS batch_id, report_start_date, report_end_date, "
        "MAX(import_frequency) AS import_frequency, COUNT(*) AS row_count, "
        "SUM(spend_usd) AS total_spend_usd, SUM(purchase_value_usd) AS total_purchase_value_usd "
        "FROM meta_ad_campaign_metrics "
        "GROUP BY report_start_date, report_end_date "
        "ORDER BY report_end_date DESC, report_start_date DESC"
    )


def _resolve_meta_ad_period(
    batch_id: int | None = None,
    start_date: str | None = None,
    end_date: str | None = None,
) -> tuple[date | None, date | None, int | None]:
    if batch_id:
        row = query_one(
            "SELECT id, report_start_date, report_end_date FROM meta_ad_import_batches WHERE id=%s",
            (batch_id,),
        )
        if not row:
            return None, None, None
        return row.get("report_start_date"), row.get("report_end_date"), row.get("id")
    if start_date and end_date:
        return _parse_meta_date(start_date), _parse_meta_date(end_date), None
    latest = query_one(
        "SELECT MAX(import_batch_id) AS batch_id, report_start_date, report_end_date "
        "FROM meta_ad_campaign_metrics "
        "GROUP BY report_start_date, report_end_date "
        "ORDER BY report_end_date DESC, report_start_date DESC LIMIT 1"
    )
    if not latest:
        return None, None, None
    return latest.get("report_start_date"), latest.get("report_end_date"), latest.get("batch_id")


_META_AD_SUMMARY_NUMERIC_FIELDS = (
    "result_count",
    "spend_usd",
    "purchase_value_usd",
    "link_clicks",
    "add_to_cart_count",
    "initiate_checkout_count",
    "impressions",
)


def _coerce_meta_product_id(value: Any) -> int | None:
    if value in (None, ""):
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _aggregate_meta_ad_summary_rows(metric_rows: list[dict]) -> list[dict]:
    grouped: dict[tuple[str, int | str], dict] = {}

    for metric in metric_rows:
        product_id = _coerce_meta_product_id(metric.get("product_id"))
        campaign_name = (metric.get("campaign_name") or "").strip()
        if product_id is not None:
            group_key: tuple[str, int | str] = ("product", product_id)
            display_name = metric.get("product_name") or campaign_name
            product_code = metric.get("matched_product_code") or metric.get("media_product_code")
        else:
            group_key = ("campaign", campaign_name or str(metric.get("id") or ""))
            display_name = campaign_name
            product_code = None

        if group_key not in grouped:
            grouped[group_key] = {
                "product_id": product_id,
                "display_name": display_name,
                "product_code": product_code,
                "campaign_count": 0,
                "_campaign_names": [],
                **{field: 0 for field in _META_AD_SUMMARY_NUMERIC_FIELDS},
            }

        row = grouped[group_key]
        if product_id is not None:
            row["display_name"] = row.get("display_name") or display_name
            row["product_code"] = row.get("product_code") or product_code
        row["campaign_count"] += 1
        if campaign_name:
            row["_campaign_names"].append(campaign_name)
        for field in _META_AD_SUMMARY_NUMERIC_FIELDS:
            row[field] += metric.get(field) or 0

    rows: list[dict] = []
    for row in grouped.values():
        campaign_names = sorted(row.pop("_campaign_names"))
        row["campaign_names"] = ", ".join(campaign_names)
        spend = row.get("spend_usd") or 0
        purchase_value = row.get("purchase_value_usd") or 0
        result_count = row.get("result_count") or 0
        row["roas_purchase"] = purchase_value / spend if spend else None
        row["cost_per_result_usd"] = spend / result_count if result_count else None
        rows.append(row)

    rows.sort(key=lambda row: row.get("spend_usd") or 0, reverse=True)
    return rows


def get_meta_ad_summary(
    batch_id: int | None = None,
    start_date: str | None = None,
    end_date: str | None = None,
) -> dict:
    report_start, report_end, resolved_batch_id = _resolve_meta_ad_period(batch_id, start_date, end_date)
    if not report_start or not report_end:
        return {"period": None, "rows": [], "unmatched": []}

    metric_rows = query(
        "SELECT m.id, m.product_id, mp.name AS product_name, mp.product_code AS media_product_code, "
        "m.matched_product_code, m.campaign_name, m.result_count, m.spend_usd, "
        "m.purchase_value_usd, m.link_clicks, m.add_to_cart_count, "
        "m.initiate_checkout_count, m.impressions "
        "FROM meta_ad_campaign_metrics m "
        "LEFT JOIN media_products mp ON mp.id = m.product_id "
        "WHERE m.report_start_date=%s AND m.report_end_date=%s "
        "ORDER BY m.spend_usd DESC",
        (report_start, report_end),
    )
    rows = _aggregate_meta_ad_summary_rows(metric_rows)

    product_ids = [int(row["product_id"]) for row in rows if row.get("product_id")]
    orders_by_product: dict[int, dict] = {}
    if product_ids:
        placeholders = ",".join(["%s"] * len(product_ids))
        order_rows = query(
            "SELECT product_id, COUNT(DISTINCT shopify_order_id) AS shopify_order_count, "
            "SUM(lineitem_quantity) AS shopify_quantity, "
            "SUM(COALESCE(lineitem_price, 0) * lineitem_quantity) AS shopify_revenue "
            "FROM shopify_orders "
            f"WHERE product_id IN ({placeholders}) AND created_at_order >= %s AND created_at_order < DATE_ADD(%s, INTERVAL 1 DAY) "
            "GROUP BY product_id",
            tuple(product_ids + [report_start, report_end]),
        )
        orders_by_product = {int(row["product_id"]): row for row in order_rows}

    for row in rows:
        order_metrics = orders_by_product.get(int(row["product_id"] or 0), {})
        row["shopify_order_count"] = order_metrics.get("shopify_order_count") or 0
        row["shopify_quantity"] = order_metrics.get("shopify_quantity") or 0
        row["shopify_revenue"] = order_metrics.get("shopify_revenue") or 0

    unmatched = query(
        "SELECT id, campaign_name, normalized_campaign_code, spend_usd, result_count, purchase_value_usd "
        "FROM meta_ad_campaign_metrics "
        "WHERE report_start_date=%s AND report_end_date=%s AND product_id IS NULL "
        "ORDER BY spend_usd DESC",
        (report_start, report_end),
    )
    return {
        "period": {
            "batch_id": resolved_batch_id,
            "report_start_date": report_start,
            "report_end_date": report_end,
        },
        "rows": rows,
        "unmatched": unmatched,
    }


# ── 导入 ───────────────────────────────────────────────

def import_orders(rows: list[dict]) -> dict:
    """将原始行批量写入 shopify_orders，去重。返回 {imported, skipped}。"""
    if not rows:
        return {"imported": 0, "skipped": 0}

    BATCH = 500
    imported = 0
    skipped = 0

    for i in range(0, len(rows), BATCH):
        batch = rows[i:i + BATCH]
        values = []
        for row in batch:
            oid_str = (row.get("Id") or "").strip()
            if not oid_str:
                skipped += 1
                continue
            try:
                oid = int(oid_str)
            except ValueError:
                skipped += 1
                continue

            name = (row.get("Lineitem name") or "").strip()
            if not name:
                skipped += 1
                continue

            values.append((
                oid,
                (row.get("Name") or "").strip()[:32] or None,
                _parse_shopify_ts(row.get("Created at", "")),
                name[:500],
                (row.get("Lineitem sku") or "").strip()[:128] or None,
                _safe_int(row.get("Lineitem quantity", "1"), 1),
                _safe_float(row.get("Lineitem price")),
                (row.get("Billing Country") or "").strip()[:8] or None,
                _safe_float(row.get("Total")),
                _safe_float(row.get("Subtotal")),
                _safe_float(row.get("Shipping")),
                (row.get("Currency") or "").strip()[:8] or None,
                (row.get("Financial Status") or "").strip()[:32] or None,
                (row.get("Fulfillment Status") or "").strip()[:32] or None,
                (row.get("Vendor") or "").strip()[:128] or None,
            ))

        if not values:
            continue

        placeholders = ",".join(["(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)"] * len(values))
        flat = []
        for v in values:
            flat.extend(v)

        sql = (
            "INSERT IGNORE INTO shopify_orders "
            "(shopify_order_id, order_name, created_at_order, lineitem_name, "
            "lineitem_sku, lineitem_quantity, lineitem_price, billing_country, "
            "total, subtotal, shipping, currency, financial_status, "
            "fulfillment_status, vendor) VALUES " + placeholders
        )
        conn = get_conn()
        try:
            with conn.cursor() as cur:
                cur.execute(sql, tuple(flat))
                affected = cur.rowcount
        finally:
            conn.close()
        imported += affected
        skipped += len(values) - affected

    return {"imported": imported, "skipped": skipped}


def get_import_stats() -> dict:
    """返回数据库统计概览。"""
    row = query_one(
        "SELECT COUNT(*) AS total_rows, "
        "MIN(created_at_order) AS min_date, MAX(created_at_order) AS max_date, "
        "COUNT(DISTINCT lineitem_name) AS product_count, "
        "COUNT(DISTINCT billing_country) AS country_count, "
        "SUM(CASE WHEN product_id IS NOT NULL THEN 1 ELSE 0 END) AS matched_rows "
        "FROM shopify_orders"
    )
    return row or {}


# ── 产品标题缓存 ──────────────────────────────────────

def fetch_product_page_title(product_code: str) -> str | None:
    """从英文产品页抓取 <title>。"""
    url = f"https://newjoyloo.com/products/{product_code}"
    try:
        resp = requests.get(url, timeout=8, allow_redirects=True,
                            headers={"User-Agent": "Mozilla/5.0 AutoVideoSrt/1.0"})
        if resp.status_code != 200:
            return None
        m = _TITLE_RE.search(resp.text)
        if m:
            title = m.group(1).strip()
            # 先解码 HTML 实体
            title = (title
                     .replace("&ndash;", "–")
                     .replace("&mdash;", "—")
                     .replace("&amp;", "&")
                     .replace("&lt;", "<")
                     .replace("&gt;", ">")
                     .replace("&quot;", '"'))
            # 再去掉 " | Store Name" / " – Store Name" / " - Store Name" 后缀
            for sep in (" | ", " – ", " — ", " - "):
                if sep in title:
                    title = title.rsplit(sep, 1)[0].strip()
                    break
            return title[:500] if title else None
    except requests.RequestException as exc:
        log.debug("fetch title failed for %s: %s", product_code, exc)
    return None


def refresh_product_titles(product_ids: list[int] | None = None) -> dict:
    """批量刷新产品标题缓存。返回 {fetched, skipped, errors}。"""
    if product_ids:
        placeholders = ",".join(["%s"] * len(product_ids))
        products = query(
            f"SELECT id, product_code FROM media_products "
            f"WHERE id IN ({placeholders}) AND product_code IS NOT NULL AND deleted_at IS NULL",
            tuple(product_ids),
        )
    else:
        products = query(
            "SELECT mp.id, mp.product_code FROM media_products mp "
            "LEFT JOIN product_title_cache ptc ON ptc.product_id = mp.id "
            "WHERE mp.product_code IS NOT NULL AND mp.deleted_at IS NULL "
            "AND (ptc.id IS NULL OR ptc.fetched_at < DATE_SUB(NOW(), INTERVAL 7 DAY))",
        )

    fetched = 0
    errors = 0
    for p in products:
        code = p["product_code"]
        if not code:
            continue
        title = fetch_product_page_title(code)
        if title:
            execute(
                "INSERT INTO product_title_cache (product_id, product_code, page_title, fetched_at) "
                "VALUES (%s, %s, %s, NOW()) "
                "ON DUPLICATE KEY UPDATE page_title=VALUES(page_title), fetched_at=NOW()",
                (p["id"], code, title),
            )
            fetched += 1
        else:
            errors += 1
        time.sleep(0.5)  # 限速

    return {"fetched": fetched, "skipped": len(products) - fetched - errors, "errors": errors}


def match_orders_to_products() -> int:
    """将 lineitem_name 匹配到 product_id。返回新匹配行数。"""
    # 前缀匹配：订单商品名以产品标题开头（后面可能有变体信息如 "- 1 Pack"）
    affected = execute(
        "UPDATE shopify_orders so "
        "JOIN product_title_cache ptc ON so.lineitem_name LIKE CONCAT(ptc.page_title, '%') "
        "SET so.product_id = ptc.product_id "
        "WHERE so.product_id IS NULL"
    )
    return affected


# ── 国家 ↔ 语种映射 ───────────────────────────────────────


COUNTRY_TO_LANG: dict[str, str] = {
    "US": "en", "GB": "en", "UK": "en",
    "AU": "en", "CA": "en", "IE": "en", "NZ": "en",
    "DE": "de", "AT": "de",
    "FR": "fr",
    "ES": "es",
    "IT": "it",
    "NL": "nl",
    "SE": "sv",
    "FI": "fi",
    "JP": "ja",
    "KR": "ko",
    "BR": "pt-BR",
    "PT": "pt",
}


# 同语种多国家时，列输出顺序固定走这张表（不出现的语种按 dict 插入序）
LANG_PRIORITY_COUNTRIES: dict[str, list[str]] = {
    "en": ["US", "GB", "AU", "CA", "IE", "NZ"],
    "de": ["DE", "AT"],
}


def _load_enabled_lang_codes() -> list[str]:
    """读取 media_languages.enabled=1 的语种 code，按 sort_order 升序。

    与 appcore.medias.list_enabled_language_codes() 等价；放在本模块里独立维护，
    便于单测通过 monkeypatch.setattr(oa, "_load_enabled_lang_codes", …) 替换实现，
    而不必污染 appcore.medias。
    """
    rows = query(
        "SELECT code FROM media_languages "
        "WHERE enabled=1 ORDER BY sort_order ASC, code ASC"
    )
    return [r["code"] for r in rows]


def get_enabled_country_columns() -> list[dict]:
    """根据 media_languages 启用语种推导出"国家列"序列。

    返回列表如 [{"country": "US", "lang": "en"}, …]，按
    sort_order(语种) → LANG_PRIORITY_COUNTRIES(同语种内部顺序) 双重排序。
    未在 COUNTRY_TO_LANG 里出现的启用语种被静默跳过（不报错）。
    """
    enabled_langs = _load_enabled_lang_codes()
    columns: list[dict] = []
    seen: set[str] = set()

    # 反向构建：lang → [country, ...]，对未在优先表里的语种走 dict 插入序
    lang_to_countries: dict[str, list[str]] = {}
    for country, lang in COUNTRY_TO_LANG.items():
        lang_to_countries.setdefault(lang, []).append(country)
    # 优先表覆盖默认顺序
    for lang, ordered in LANG_PRIORITY_COUNTRIES.items():
        if lang in lang_to_countries:
            lang_to_countries[lang] = ordered

    for lang in enabled_langs:
        countries = lang_to_countries.get(lang)
        if not countries:
            continue
        for country in countries:
            if country in seen:
                continue
            seen.add(country)
            columns.append({"country": country, "lang": lang})

    return columns


# ── 分析查询 ───────────────────────────────────────────

def _month_range(year: int, month: int) -> tuple[str, str]:
    """返回 (start, end) 字符串用于 WHERE created_at_order >= start AND < end。"""
    start = f"{year:04d}-{month:02d}-01"
    if month == 12:
        end = f"{year + 1:04d}-01-01"
    else:
        end = f"{year:04d}-{month + 1:02d}-01"
    return start, end


def get_monthly_summary(year: int, month: int, product_id: int | None = None) -> dict:
    """月度汇总：按产品 × 国家。"""
    start, end = _month_range(year, month)
    extra_filter = ""
    args: list[Any] = [start, end]
    if product_id is not None:
        extra_filter = "AND so.product_id = %s"
        args.append(product_id)

    # 按产品汇总
    products = query(
        f"SELECT so.product_id, "
        f"COALESCE(mp.name, ptc.page_title, so.lineitem_name) AS display_name, "
        f"mp.product_code, "
        f"SUM(so.lineitem_quantity) AS total_qty, "
        f"COUNT(DISTINCT so.shopify_order_id) AS order_count, "
        f"SUM(COALESCE(so.lineitem_price,0) * so.lineitem_quantity) AS total_revenue "
        f"FROM shopify_orders so "
        f"LEFT JOIN product_title_cache ptc ON ptc.product_id = so.product_id "
        f"LEFT JOIN media_products mp ON mp.id = so.product_id "
        f"WHERE so.created_at_order >= %s AND so.created_at_order < %s {extra_filter} "
        f"GROUP BY so.product_id, display_name, mp.product_code "
        f"ORDER BY total_qty DESC",
        tuple(args),
    )

    # 按国家汇总
    countries = query(
        f"SELECT billing_country, "
        f"SUM(lineitem_quantity) AS total_qty, "
        f"COUNT(DISTINCT shopify_order_id) AS order_count "
        f"FROM shopify_orders "
        f"WHERE created_at_order >= %s AND created_at_order < %s {extra_filter} "
        f"GROUP BY billing_country ORDER BY total_qty DESC",
        tuple(args),
    )

    # 产品 × 国家矩阵
    matrix_rows = query(
        f"SELECT so.product_id, "
        f"COALESCE(mp.name, ptc.page_title, so.lineitem_name) AS display_name, "
        f"so.billing_country, "
        f"SUM(so.lineitem_quantity) AS total_qty "
        f"FROM shopify_orders so "
        f"LEFT JOIN product_title_cache ptc ON ptc.product_id = so.product_id "
        f"LEFT JOIN media_products mp ON mp.id = so.product_id "
        f"WHERE so.created_at_order >= %s AND so.created_at_order < %s {extra_filter} "
        f"GROUP BY so.product_id, display_name, so.billing_country "
        f"ORDER BY display_name, total_qty DESC",
        tuple(args),
    )

    # 组装矩阵
    country_list = [c["billing_country"] or "未知" for c in countries]
    matrix: dict[str, dict[str, int]] = {}
    product_order: list[str] = []
    for mr in matrix_rows:
        dn = mr["display_name"] or "未知"
        if dn not in matrix:
            matrix[dn] = {}
            product_order.append(dn)
        matrix[dn][mr["billing_country"] or "未知"] = mr["total_qty"]

    # 素材数量：按 product × lang，复用 dashboard 已有的统计逻辑
    media_counts_all = _count_media_items_by_product()
    if product_id is not None:
        media_counts = (
            {product_id: media_counts_all[product_id]}
            if product_id in media_counts_all
            else {}
        )
    else:
        # 仅保留本次查询里出现的产品，避免响应膨胀
        active_pids = {p["product_id"] for p in products if p.get("product_id") is not None}
        media_counts = {
            pid: counts for pid, counts in media_counts_all.items() if pid in active_pids
        }

    country_columns = get_enabled_country_columns()

    return {
        "products": products,
        "countries": countries,
        "country_list": country_list,
        "matrix": matrix,
        "product_order": product_order,
        "country_columns": country_columns,
        "media_counts": media_counts,
    }


def get_product_country_detail(product_id: int, year: int, month: int) -> list[dict]:
    """单个产品在指定月份的"国家×素材×订单"明细。

    覆盖所有启用国家，即使该国当月 0 单 0 素材，也会输出一行（值全 0）。

    返回每行字段：country / lang / qty / orders / revenue / media_count
    """
    start, end = _month_range(year, month)

    # 该产品在月份内的国家汇总
    rows = query(
        "SELECT so.billing_country, "
        "SUM(so.lineitem_quantity) AS qty, "
        "COUNT(DISTINCT so.shopify_order_id) AS orders, "
        "SUM(COALESCE(so.lineitem_price, 0) * so.lineitem_quantity) AS revenue "
        "FROM shopify_orders so "
        "WHERE so.product_id = %s "
        "AND so.created_at_order >= %s AND so.created_at_order < %s "
        "GROUP BY so.billing_country",
        (product_id, start, end),
    )
    by_country: dict[str, dict] = {}
    for r in rows:
        country = r.get("billing_country") or ""
        by_country[country] = {
            "qty": int(r.get("qty") or 0),
            "orders": int(r.get("orders") or 0),
            "revenue": float(r.get("revenue") or 0),
        }

    # 该产品的素材语种分布
    media_rows = query(
        "SELECT lang, COUNT(*) AS n FROM media_items "
        "WHERE product_id = %s AND deleted_at IS NULL "
        "GROUP BY lang",
        (product_id,),
    )
    media_by_lang: dict[str, int] = {}
    for r in media_rows:
        lang = r.get("lang") or ""
        media_by_lang[lang] = int(r.get("n") or 0)

    out: list[dict] = []
    for col in get_enabled_country_columns():
        country = col["country"]
        lang = col["lang"]
        order_data = by_country.get(country, {})
        out.append({
            "country": country,
            "lang": lang,
            "qty": order_data.get("qty", 0),
            "orders": order_data.get("orders", 0),
            "revenue": round(order_data.get("revenue", 0.0), 2),
            "media_count": media_by_lang.get(lang, 0),
        })
    return out


def get_daily_detail(year: int, month: int, product_id: int | None = None) -> list[dict]:
    """每日明细：按日期 × 产品 × 国家。"""
    start, end = _month_range(year, month)
    extra_filter = ""
    args: list[Any] = [start, end]
    if product_id is not None:
        extra_filter = "AND so.product_id = %s"
        args.append(product_id)

    return query(
        f"SELECT DATE(so.created_at_order) AS sale_date, "
        f"so.product_id, "
        f"COALESCE(mp.name, ptc.page_title, so.lineitem_name) AS display_name, "
        f"so.billing_country, "
        f"SUM(so.lineitem_quantity) AS total_qty, "
        f"COUNT(DISTINCT so.shopify_order_id) AS order_count "
        f"FROM shopify_orders so "
        f"LEFT JOIN product_title_cache ptc ON ptc.product_id = so.product_id "
        f"LEFT JOIN media_products mp ON mp.id = so.product_id "
        f"WHERE so.created_at_order >= %s AND so.created_at_order < %s {extra_filter} "
        f"GROUP BY sale_date, so.product_id, display_name, so.billing_country "
        f"ORDER BY sale_date ASC, total_qty DESC",
        tuple(args),
    )


def get_weekly_summary(year: int, week: int) -> dict:
    """周汇总：按 ISO 周。"""
    target = f"{year:04d}{week:02d}"
    products = query(
        "SELECT so.product_id, "
        "COALESCE(mp.name, ptc.page_title, so.lineitem_name) AS display_name, "
        "SUM(so.lineitem_quantity) AS total_qty, "
        "COUNT(DISTINCT so.shopify_order_id) AS order_count "
        "FROM shopify_orders so "
        "LEFT JOIN product_title_cache ptc ON ptc.product_id = so.product_id "
        "LEFT JOIN media_products mp ON mp.id = so.product_id "
        "WHERE YEARWEEK(so.created_at_order, 1) = %s "
        "GROUP BY so.product_id, display_name ORDER BY total_qty DESC",
        (target,),
    )
    countries = query(
        "SELECT billing_country, SUM(lineitem_quantity) AS total_qty "
        "FROM shopify_orders "
        "WHERE YEARWEEK(created_at_order, 1) = %s "
        "GROUP BY billing_country ORDER BY total_qty DESC",
        (target,),
    )
    return {"products": products, "countries": countries}


def search_products(q: str) -> list[dict]:
    """按产品 ID 或标题搜索。"""
    like = f"%{q}%"
    # 尝试将 q 解析为数字（product_id）
    try:
        pid = int(q)
    except ValueError:
        pid = None

    if pid is not None:
        return query(
            "SELECT DISTINCT so.product_id, "
            "COALESCE(mp.name, ptc.page_title, so.lineitem_name) AS display_name, "
            "mp.product_code "
            "FROM shopify_orders so "
            "LEFT JOIN product_title_cache ptc ON ptc.product_id = so.product_id "
            "LEFT JOIN media_products mp ON mp.id = so.product_id "
            "WHERE so.product_id = %s OR so.lineitem_name LIKE %s "
            "LIMIT 50",
            (pid, like),
        )
    return query(
        "SELECT DISTINCT so.product_id, "
        "COALESCE(mp.name, ptc.page_title, so.lineitem_name) AS display_name, "
        "mp.product_code "
        "FROM shopify_orders so "
        "LEFT JOIN product_title_cache ptc ON ptc.product_id = so.product_id "
        "LEFT JOIN media_products mp ON mp.id = so.product_id "
        "WHERE so.lineitem_name LIKE %s OR ptc.page_title LIKE %s "
        "LIMIT 50",
        (like, like),
    )


def get_available_months() -> list[dict]:
    """返回有数据的年月列表。"""
    return query(
        "SELECT YEAR(created_at_order) AS y, MONTH(created_at_order) AS m, "
        "COUNT(*) AS row_count "
        "FROM shopify_orders "
        "GROUP BY YEAR(created_at_order), MONTH(created_at_order) "
        "ORDER BY y DESC, m DESC"
    )


# ── 产品看板 V1 ───────────────────────────────────────────

def _compute_pct_change(now, prev) -> float | None:
    """环比百分比。返回 None 表示无法计算（prev=0 且 now>0）。"""
    now_v = float(now or 0)
    prev_v = float(prev or 0)
    if prev_v == 0 and now_v == 0:
        return 0.0
    if prev_v == 0:
        return None
    return round((now_v - prev_v) / prev_v * 100, 2)


def _resolve_period_range(
    period: str,
    *,
    year: int | None = None,
    month: int | None = None,
    week: int | None = None,
    date_str: str | None = None,
    today: date | None = None,
) -> tuple[date, date]:
    """返回 (start, end) 闭区间。

    - month: 该月 1 日 ~ 月末；若为当月，end = 昨日（不含今天）
    - week: ISO 周一 ~ 周日；若为当周，end = 昨日
    - day: date_str ~ date_str
    """
    today = today or date.today()
    yesterday = today - timedelta(days=1)

    if period == "month":
        if year is None or month is None:
            raise ValueError("year and month required for period=month")
        start = date(year, month, 1)
        last_day = calendar.monthrange(year, month)[1]
        end = date(year, month, last_day)
        if start <= today <= end:
            end = yesterday if yesterday >= start else start
        return start, end

    if period == "week":
        if year is None or week is None:
            raise ValueError("year and week required for period=week")
        # ISO week: %G-%V-%u; %u=1 = Monday
        start = datetime.strptime(f"{year}-{week:02d}-1", "%G-%V-%u").date()
        end = start + timedelta(days=6)
        if start <= today <= end:
            end = yesterday if yesterday >= start else start
        return start, end

    if period == "day":
        if not date_str:
            raise ValueError("date required for period=day")
        d = datetime.strptime(date_str, "%Y-%m-%d").date()
        return d, d

    raise ValueError(f"invalid period: {period}")


def _resolve_compare_range(start: date, end: date, period: str) -> tuple[date, date]:
    """返回上一个同长度切片。月模式下，若上月天数较少则将日期 clamp 到上月末日。"""
    if period == "month":
        # 减一个月：直接调整 month 字段
        prev_year = start.year - (1 if start.month == 1 else 0)
        prev_month = 12 if start.month == 1 else start.month - 1
        prev_month_last = calendar.monthrange(prev_year, prev_month)[1]
        prev_start = date(prev_year, prev_month, min(start.day, prev_month_last))
        # end 取上月同一天（截断到上月末尾）
        prev_end = date(prev_year, prev_month, min(end.day, prev_month_last))
        return prev_start, prev_end

    if period == "week":
        prev_start = start - timedelta(days=7)
        return prev_start, prev_start + (end - start)

    if period == "day":
        prev = start - timedelta(days=1)
        return prev, prev

    raise ValueError(f"invalid period: {period}")


def _aggregate_orders_by_product(
    start: date, end: date, *, country: str | None = None
) -> dict[int, dict]:
    """按产品聚合订单。返回 {product_id: {orders, units, revenue}}。"""
    sql = (
        "SELECT product_id, "
        "COUNT(DISTINCT shopify_order_id) AS orders, "
        "SUM(lineitem_quantity) AS units, "
        "SUM(COALESCE(lineitem_price, 0) * lineitem_quantity) AS revenue "
        "FROM shopify_orders "
        "WHERE created_at_order >= %s AND created_at_order < DATE_ADD(%s, INTERVAL 1 DAY) "
    )
    args: tuple = (start, end)
    if country:
        sql += "AND billing_country = %s "
        args = (start, end, country)
    sql += "GROUP BY product_id"

    rows = query(sql, args)
    out: dict[int, dict] = {}
    for r in rows:
        pid = r.get("product_id")
        if pid is None:
            continue
        out[int(pid)] = {
            "orders": int(r.get("orders") or 0),
            "units": int(r.get("units") or 0),
            "revenue": float(r.get("revenue") or 0),
        }
    return out


def _aggregate_ads_by_product(start: date, end: date) -> dict[int, dict]:
    """按产品聚合广告。仅纳入 [report_start_date, report_end_date] 完全
    被 [start, end] 覆盖的报表（决策 #7）。
    返回 {product_id: {spend, purchases, purchase_value}}。"""
    sql = (
        "SELECT product_id, "
        "SUM(spend_usd) AS spend, "
        "SUM(result_count) AS purchases, "
        "SUM(purchase_value_usd) AS purchase_value "
        "FROM meta_ad_campaign_metrics "
        "WHERE report_start_date >= %s AND report_end_date <= %s "
        "GROUP BY product_id"
    )
    rows = query(sql, (start, end))
    out: dict[int, dict] = {}
    for r in rows:
        pid = r.get("product_id")
        if pid is None:
            continue
        out[int(pid)] = {
            "spend": float(r.get("spend") or 0),
            "purchases": int(r.get("purchases") or 0),
            "purchase_value": float(r.get("purchase_value") or 0),
        }
    return out


def _count_media_items_by_product() -> dict[int, dict[str, int]]:
    """SELECT product_id, lang, COUNT(*) FROM media_items WHERE deleted_at IS NULL
       GROUP BY product_id, lang"""
    rows = query(
        "SELECT product_id, lang, COUNT(*) AS n FROM media_items "
        "WHERE deleted_at IS NULL "
        "GROUP BY product_id, lang"
    )
    out: dict[int, dict[str, int]] = {}
    for r in rows:
        pid = r.get("product_id")
        if pid is None:
            continue
        out.setdefault(int(pid), {})[r.get("lang") or ""] = int(r.get("n") or 0)
    return out


def _join_and_compute_dashboard_rows(
    *,
    products: dict[int, dict],
    orders_now: dict[int, dict],
    orders_prev: dict[int, dict],
    ads_now: dict[int, dict],
    ads_prev: dict[int, dict],
    items: dict[int, dict[str, int]],
    ad_data_available: bool,
) -> list[dict]:
    """合并 4 个数据源 + 媒体素材数 + 计算 ROAS / 环比百分比。
    决策 #12 剔除两边都 0 的产品。"""
    rows: list[dict] = []
    candidate_ids = set(orders_now.keys()) | set(ads_now.keys())
    for pid in candidate_ids:
        if pid not in products:
            # 产品已被删除/归档，跳过
            continue
        prod = products[pid]
        o_now = orders_now.get(pid, {})
        o_prev = orders_prev.get(pid, {})
        a_now = ads_now.get(pid, {})
        a_prev = ads_prev.get(pid, {})

        orders = int(o_now.get("orders") or 0)
        spend = float(a_now.get("spend") or 0)
        if orders == 0 and spend == 0:
            continue  # 决策 #12

        revenue = float(o_now.get("revenue") or 0)
        revenue_prev = float(o_prev.get("revenue") or 0)
        spend_prev = float(a_prev.get("spend") or 0)
        roas = (revenue / spend) if spend > 0 else None
        roas_prev = (revenue_prev / spend_prev) if spend_prev > 0 else None

        row = {
            "product_id": pid,
            "product_code": prod.get("product_code"),
            "product_name": prod.get("name"),
            "orders": orders,
            "orders_prev": int(o_prev.get("orders") or 0),
            "orders_pct": _compute_pct_change(orders, o_prev.get("orders")),
            "units": int(o_now.get("units") or 0),
            "units_prev": int(o_prev.get("units") or 0),
            "units_pct": _compute_pct_change(o_now.get("units"), o_prev.get("units")),
            "revenue": round(revenue, 2),
            "revenue_prev": round(revenue_prev, 2),
            "revenue_pct": _compute_pct_change(revenue, revenue_prev),
            "media_items_by_lang": items.get(pid, {}),
            "ad_data_available": ad_data_available,
        }
        if ad_data_available:
            row.update({
                "spend": round(spend, 2),
                "spend_prev": round(spend_prev, 2),
                "spend_pct": _compute_pct_change(spend, spend_prev),
                "meta_purchases": int(a_now.get("purchases") or 0),
                "meta_purchases_prev": int(a_prev.get("purchases") or 0),
                "meta_purchases_pct": _compute_pct_change(
                    a_now.get("purchases"), a_prev.get("purchases")
                ),
                "roas": round(roas, 2) if roas is not None else None,
                "roas_prev": round(roas_prev, 2) if roas_prev is not None else None,
                "roas_pct": _compute_pct_change(roas, roas_prev),
            })
        else:
            row.update({
                "spend": None, "spend_prev": None, "spend_pct": None,
                "meta_purchases": None, "meta_purchases_prev": None, "meta_purchases_pct": None,
                "roas": None, "roas_prev": None, "roas_pct": None,
            })
        rows.append(row)
    return rows


_DASHBOARD_SORT_FIELDS = frozenset({"spend", "revenue", "orders", "units", "roas"})


def get_dashboard(
    *,
    period: str,
    year: int | None = None,
    month: int | None = None,
    week: int | None = None,
    date_str: str | None = None,
    country: str | None = None,
    sort_by: str | None = None,
    sort_dir: str = "desc",
    compare: bool = True,
    search: str | None = None,
    today: date | None = None,
) -> dict:
    """产品看板查询主入口。详见 spec。"""
    today = today or date.today()
    start, end = _resolve_period_range(
        period, year=year, month=month, week=week, date_str=date_str, today=today
    )

    # 周/月支持广告；日视图不查广告（决策 #3）
    # 国家筛选启用时广告整列降级（meta_ad 表无 country 字段）
    ad_data_available = period in ("week", "month") and not country

    orders_now = _aggregate_orders_by_product(start, end, country=country)
    ads_now = _aggregate_ads_by_product(start, end) if ad_data_available else {}

    orders_prev: dict[int, dict] = {}
    ads_prev: dict[int, dict] = {}
    compare_period = None
    if compare:
        prev_start, prev_end = _resolve_compare_range(start, end, period)
        orders_prev = _aggregate_orders_by_product(prev_start, prev_end, country=country)
        ads_prev = _aggregate_ads_by_product(prev_start, prev_end) if ad_data_available else {}
        compare_period = {
            "start": prev_start.isoformat(),
            "end": prev_end.isoformat(),
            "label": _format_period_label(prev_start, prev_end, period),
        }

    items = _count_media_items_by_product()

    candidate_ids = set(orders_now.keys()) | set(ads_now.keys())
    products = _load_products(candidate_ids, search=search)

    rows = _join_and_compute_dashboard_rows(
        products=products,
        orders_now=orders_now, orders_prev=orders_prev,
        ads_now=ads_now, ads_prev=ads_prev,
        items=items,
        ad_data_available=ad_data_available,
    )

    # 排序
    sort_key = sort_by if sort_by in _DASHBOARD_SORT_FIELDS else (
        "spend" if ad_data_available else "revenue"
    )
    reverse = (sort_dir.lower() == "desc")
    rows.sort(key=lambda r: (r.get(sort_key) is None, r.get(sort_key) or 0), reverse=reverse)

    summary = _summarize_dashboard(rows, ad_data_available)

    return {
        "period": {
            "start": start.isoformat(),
            "end": end.isoformat(),
            "label": _format_period_label(start, end, period),
        },
        "compare_period": compare_period,
        "country": country,
        "products": rows,
        "summary": summary,
    }


def _format_period_label(start: date, end: date, period: str) -> str:
    if period == "month":
        if start.day == 1 and end.day == calendar.monthrange(start.year, start.month)[1]:
            return f"{start.year} 年 {start.month} 月"
        return f"{start.year} 年 {start.month} 月（{start.day}-{end.day} 日）"
    if period == "week":
        return f"{start.isoformat()} ~ {end.isoformat()}"
    return start.isoformat()


def _load_products(ids: set[int], *, search: str | None = None) -> dict[int, dict]:
    """查询产品基础信息。
    始终过滤 archived/deleted；search 启用时附加 name/product_code LIKE 过滤；
    始终用 ids IN 限制为本期有数据的产品（避免无活动产品出现在看板上）。"""
    if not ids:
        return {}
    placeholders = ", ".join(["%s"] * len(ids))
    sql = (
        f"SELECT id, name, product_code FROM media_products "
        f"WHERE id IN ({placeholders}) "
        f"AND (archived = 0 OR archived IS NULL) AND deleted_at IS NULL"
    )
    args: tuple = tuple(ids)
    if search:
        like = f"%{search}%"
        sql += " AND (name LIKE %s OR product_code LIKE %s)"
        args = args + (like, like)
    rows = query(sql, args)
    return {int(r["id"]): r for r in rows}


def _summarize_dashboard(rows: list[dict], ad_data_available: bool) -> dict:
    total_orders = sum(r.get("orders") or 0 for r in rows)
    total_revenue = round(sum(r.get("revenue") or 0 for r in rows), 2)
    summary = {
        "total_orders": total_orders,
        "total_revenue": total_revenue,
    }
    if ad_data_available:
        total_spend = round(sum(r.get("spend") or 0 for r in rows), 2)
        summary["total_spend"] = total_spend
        summary["total_meta_purchases"] = sum(r.get("meta_purchases") or 0 for r in rows)
        summary["total_roas"] = round(total_revenue / total_spend, 2) if total_spend > 0 else None
    else:
        summary["total_spend"] = None
        summary["total_meta_purchases"] = None
        summary["total_roas"] = None
    return summary
